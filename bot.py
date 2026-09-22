#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
بوت تيليجرام لتسجيل أيام الشغل (Gig Tracker) - النسخة المطوّرة
Made by OmarMidlg

مزايا هاي النسخة:
- أزرار تفاعلية (Inline) بدل الكتابة اليدوية
- تعديل أي يومية بعد إضافتها
- بحث/فلترة الأيام حسب التاريخ / المكان / نوع الشغل / الشريك
- تصدير البيانات Excel أو PDF (يدعم العربي)
- تقرير شهري/سنوي مع رسم بياني
- تذكير أسبوعي تلقائي بملخص الحساب
"""

import calendar
import io
import logging
import os
import sqlite3
from contextlib import closing
from datetime import datetime, date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

import arabic_reshaper
from bidi.algorithm import get_display

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from fpdf import FPDF

from telegram import (
    ReplyKeyboardMarkup,
    KeyboardButton,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    Update,
)
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)

# ----------------------------------------------------------------------
# إعدادات عامة
# ----------------------------------------------------------------------

BOT_TOKEN = os.environ.get("BOT_TOKEN", "8789370643:AAFlsbbu3GCllvundSrCPqBpGe5o08VD4RY")
DB_PATH = os.environ.get("DB_PATH", "gig_tracker.db")
DEFAULT_FEE = 120.00  # الأجرة الافتراضية لكل حدث بالشيكل (NIS)
PAGE_SIZE = 5

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ARABIC_FONT_PATH = os.path.join(BASE_DIR, "NotoNaskhArabic-Regular.ttf")

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)

# تسجيل الخط العربي بمكتبة matplotlib (للرسوم البيانية)
if os.path.exists(ARABIC_FONT_PATH):
    font_manager.fontManager.addfont(ARABIC_FONT_PATH)
    ARABIC_FONT_NAME = font_manager.FontProperties(fname=ARABIC_FONT_PATH).get_name()
else:
    ARABIC_FONT_NAME = None
    logger.warning("ملف الخط العربي غير موجود: %s - الرسوم البيانية والـPDF قد لا تدعم العربي.", ARABIC_FONT_PATH)

ARABIC_DAYS = {
    6: "الأحد",     # Monday=0 ... Sunday=6 في بايثون
    0: "الإثنين",
    1: "الثلاثاء",
    2: "الأربعاء",
    3: "الخميس",
    4: "الجمعة",
    5: "السبت",
}
ARABIC_DAY_ORDER = [6, 0, 1, 2, 3, 4, 5]  # الأحد أول يوم بالأسبوع عادةً

ARABIC_MONTHS = {
    1: "يناير", 2: "فبراير", 3: "مارس", 4: "إبريل", 5: "مايو", 6: "يونيو",
    7: "يوليو", 8: "أغسطس", 9: "سبتمبر", 10: "أكتوبر", 11: "نوفمبر", 12: "ديسمبر",
}

FIELD_CODES = {"d": "event_date", "l": "location", "j": "job", "p": "partner", "m": "money"}
FIELD_LABELS_AR = {
    "d": "📅 التاريخ",
    "l": "📍 المكان",
    "j": "🛠️ نوع الشغل",
    "p": "👤 الشريك",
    "m": "💵 المبلغ",
}

SETTINGS_FIELDS = {
    "bal": ("previous_balance", "الرصيد من الحساب السابق"),
    "disc": ("discount", "مبلغ الخصم"),
    "rec": ("amount_received", "المبلغ المستلم"),
    "fee": ("default_fee", "الأجرة الافتراضية لكل حدث"),
}

CANCEL_HINT = "\n\nأرسل /cancel بأي وقت للإلغاء."

MAIN_MENU = ReplyKeyboardMarkup(
    [
        [KeyboardButton("➕ إضافة يومية")],
        [KeyboardButton("📋 آخر الأيام"), KeyboardButton("📊 ملخص الحساب")],
        [KeyboardButton("🔍 بحث"), KeyboardButton("📤 تصدير البيانات")],
        [KeyboardButton("📈 تقرير شهري/سنوي"), KeyboardButton("⏰ التذكير الأسبوعي")],
        [KeyboardButton("⚙️ الإعدادات")],
    ],
    resize_keyboard=True,
)


# ----------------------------------------------------------------------
# قاعدة البيانات
# ----------------------------------------------------------------------

def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.row_factory = sqlite3.Row
    return conn


def _add_column_if_missing(conn, table, column, ddl):
    cols = [r["name"] for r in conn.execute(f"PRAGMA table_info({table})")]
    if column not in cols:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {ddl}")


def init_db():
    with closing(get_conn()) as conn, conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                full_name TEXT,
                previous_balance REAL NOT NULL DEFAULT 0,
                discount REAL NOT NULL DEFAULT 0,
                amount_received REAL NOT NULL DEFAULT 0,
                default_fee REAL NOT NULL DEFAULT %f
            )
            """
            % DEFAULT_FEE
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                event_number INTEGER NOT NULL,
                event_date TEXT NOT NULL,
                day_name TEXT NOT NULL,
                location TEXT,
                job TEXT,
                partner TEXT,
                money REAL NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (user_id) REFERENCES users(user_id)
            )
            """
        )
        # ترقية قاعدة البيانات القديمة (تذكير أسبوعي)
        _add_column_if_missing(conn, "users", "reminder_enabled", "reminder_enabled INTEGER NOT NULL DEFAULT 0")
        _add_column_if_missing(conn, "users", "reminder_day", "reminder_day INTEGER NOT NULL DEFAULT 4")
        _add_column_if_missing(conn, "users", "reminder_hour", "reminder_hour INTEGER NOT NULL DEFAULT 18")
        _add_column_if_missing(conn, "users", "last_reminder_sent", "last_reminder_sent TEXT")


def ensure_user(user_id: int, full_name: str = ""):
    with closing(get_conn()) as conn, conn:
        conn.execute(
            """
            INSERT INTO users (user_id, full_name, previous_balance, discount,
                                amount_received, default_fee)
            VALUES (?, ?, 0, 0, 0, ?)
            ON CONFLICT(user_id) DO NOTHING
            """,
            (user_id, full_name, DEFAULT_FEE),
        )


def get_user_settings(user_id: int):
    with closing(get_conn()) as conn:
        return conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone()


def get_all_users_with_reminders():
    with closing(get_conn()) as conn:
        return conn.execute("SELECT * FROM users WHERE reminder_enabled = 1").fetchall()


def get_next_event_number(user_id: int) -> int:
    with closing(get_conn()) as conn:
        row = conn.execute(
            "SELECT COALESCE(MAX(event_number), 0) + 1 AS n FROM events WHERE user_id = ?",
            (user_id,),
        ).fetchone()
        return row["n"]


def add_event(user_id, event_number, event_date, day_name, location, job, partner, money):
    with closing(get_conn()) as conn, conn:
        conn.execute(
            """
            INSERT INTO events (user_id, event_number, event_date, day_name,
                                 location, job, partner, money, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                user_id, event_number, event_date, day_name,
                location, job, partner, money,
                datetime.now().isoformat(timespec="seconds"),
            ),
        )


def _build_filter_where(user_id, filters):
    where = ["user_id = ?"]
    params = [user_id]
    filters = filters or {}
    if filters.get("location"):
        where.append("location LIKE ?")
        params.append(f"%{filters['location']}%")
    if filters.get("job"):
        where.append("job LIKE ?")
        params.append(f"%{filters['job']}%")
    if filters.get("partner"):
        where.append("partner LIKE ?")
        params.append(f"%{filters['partner']}%")
    if filters.get("date"):
        where.append("event_date = ?")
        params.append(filters["date"])
    return " AND ".join(where), params


def count_events(user_id, filters=None):
    where, params = _build_filter_where(user_id, filters)
    with closing(get_conn()) as conn:
        row = conn.execute(f"SELECT COUNT(*) AS n FROM events WHERE {where}", params).fetchone()
        return row["n"]


def get_events_paged(user_id, page, page_size=PAGE_SIZE, filters=None):
    where, params = _build_filter_where(user_id, filters)
    offset = max(page, 0) * page_size
    with closing(get_conn()) as conn:
        return conn.execute(
            f"SELECT * FROM events WHERE {where} ORDER BY event_number DESC LIMIT ? OFFSET ?",
            params + [page_size, offset],
        ).fetchall()


def get_all_events(user_id):
    with closing(get_conn()) as conn:
        return conn.execute(
            "SELECT * FROM events WHERE user_id = ? ORDER BY event_number ASC", (user_id,)
        ).fetchall()


def get_event_by_number(user_id, event_number):
    with closing(get_conn()) as conn:
        return conn.execute(
            "SELECT * FROM events WHERE user_id = ? AND event_number = ?",
            (user_id, event_number),
        ).fetchone()


def get_event_stats(user_id):
    with closing(get_conn()) as conn:
        row = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(money), 0) AS total FROM events WHERE user_id = ?",
            (user_id,),
        ).fetchone()
        return row["n"], row["total"]


def delete_event(user_id, event_number) -> bool:
    with closing(get_conn()) as conn, conn:
        cur = conn.execute(
            "DELETE FROM events WHERE user_id = ? AND event_number = ?",
            (user_id, event_number),
        )
        return cur.rowcount > 0


def update_event_field(user_id, event_number, column, value):
    assert column in {"event_date", "day_name", "location", "job", "partner", "money"}
    with closing(get_conn()) as conn, conn:
        conn.execute(
            f"UPDATE events SET {column} = ? WHERE user_id = ? AND event_number = ?",
            (value, user_id, event_number),
        )


def update_user_field(user_id, field, value):
    assert field in {"previous_balance", "discount", "amount_received", "default_fee"}
    with closing(get_conn()) as conn, conn:
        conn.execute(f"UPDATE users SET {field} = ? WHERE user_id = ?", (value, user_id))


def set_reminder(user_id, enabled=None, day=None, hour=None):
    with closing(get_conn()) as conn, conn:
        if enabled is not None:
            conn.execute("UPDATE users SET reminder_enabled = ? WHERE user_id = ?", (int(enabled), user_id))
        if day is not None:
            conn.execute("UPDATE users SET reminder_day = ? WHERE user_id = ?", (day, user_id))
        if hour is not None:
            conn.execute("UPDATE users SET reminder_hour = ? WHERE user_id = ?", (hour, user_id))


def set_last_reminder_sent(user_id, date_str):
    with closing(get_conn()) as conn, conn:
        conn.execute("UPDATE users SET last_reminder_sent = ? WHERE user_id = ?", (date_str, user_id))


def get_monthly_totals(user_id, months_back=6):
    """بيرجع قائمة (تسمية الشهر، المجموع) لآخر عدد أشهر محدد، من الأقدم للأحدث."""
    events = get_all_events(user_id)
    buckets = {}
    for e in events:
        try:
            dt = datetime.strptime(e["event_date"], "%d/%m/%Y")
        except ValueError:
            continue
        key = (dt.year, dt.month)
        buckets[key] = buckets.get(key, 0) + e["money"]

    today = date.today()
    keys = []
    y, m = today.year, today.month
    for _ in range(months_back):
        keys.append((y, m))
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    keys.reverse()

    result = []
    for (y, m) in keys:
        label = f"{ARABIC_MONTHS[m]} {y}"
        result.append((label, buckets.get((y, m), 0)))
    return result


def get_yearly_totals(user_id):
    events = get_all_events(user_id)
    buckets = {}
    for e in events:
        try:
            dt = datetime.strptime(e["event_date"], "%d/%m/%Y")
        except ValueError:
            continue
        buckets[dt.year] = buckets.get(dt.year, 0) + e["money"]
    return sorted(buckets.items())


# ----------------------------------------------------------------------
# أدوات مساعدة عامة
# ----------------------------------------------------------------------

def fmt_money(x: float) -> str:
    return f"{x:,.2f} ₪"


def parse_date(text: str):
    text = text.strip()
    if text in ("اليوم", "اليوم.", "today"):
        return datetime.now()
    text = text.replace(".", "/").replace("-", "/")
    for fmt in ("%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def parse_number(text: str):
    text = text.strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def has_arabic(text: str) -> bool:
    return any("\u0600" <= ch <= "\u06FF" for ch in str(text))


def ar_txt(text: str) -> str:
    """بيهيّئ نص عربي للعرض الصحيح بالرسوم البيانية وملفات PDF (يمين لشمال)."""
    reshaped = arabic_reshaper.reshape(text)
    return get_display(reshaped)


def build_summary_text(user_id: int) -> str:
    settings = get_user_settings(user_id)
    n_events, total_money = get_event_stats(user_id)

    previous_balance = settings["previous_balance"]
    discount = settings["discount"]
    amount_received = settings["amount_received"]

    net_amount = total_money
    remaining_balance = previous_balance + net_amount - discount - amount_received

    return (
        "📊 *ملخص الحساب*\n"
        "━━━━━━━━━━━━━━\n"
        f"الرصيد من الحساب السابق: {fmt_money(previous_balance)}\n"
        f"عدد الأيام (الأحداث): {n_events}\n"
        f"الصافي من الأيام المسجّلة: {fmt_money(net_amount)}\n"
        f"مبلغ الخصم: {fmt_money(discount)}\n"
        f"المبلغ المستلم: {fmt_money(amount_received)}\n"
        "━━━━━━━━━━━━━━\n"
        f"💰 *الرصيد المتبقي: {fmt_money(remaining_balance)}*\n"
    )


def reset_flow(context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("awaiting", None)
    context.user_data.pop("new_event", None)


# ----------------------------------------------------------------------
# عرض قائمة الأيام مع أزرار تفاعلية وترقيم صفحات
# ----------------------------------------------------------------------

def render_events_page(user_id, page, filters=None):
    filters = filters or {}
    total = count_events(user_id, filters)
    total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
    page = max(0, min(page, total_pages - 1))
    events = get_events_paged(user_id, page, PAGE_SIZE, filters)

    header = "📋 *الأيام المسجّلة*"
    if filters:
        parts = []
        if filters.get("location"):
            parts.append(f"المكان يحتوي: {filters['location']}")
        if filters.get("job"):
            parts.append(f"الشغل يحتوي: {filters['job']}")
        if filters.get("partner"):
            parts.append(f"الشريك يحتوي: {filters['partner']}")
        if filters.get("date"):
            parts.append(f"التاريخ: {filters['date']}")
        header += f"\n🔎 نتائج البحث ({' | '.join(parts)})"

    if not events:
        text = header + "\n\nلا توجد نتائج."
        rows = []
        if filters:
            rows.append([InlineKeyboardButton("❌ مسح البحث", callback_data="lp:clear:0")])
        return text, InlineKeyboardMarkup(rows) if rows else None

    lines = [header, f"صفحة {page + 1} من {total_pages} (المجموع: {total})\n"]
    rows = []
    for e in events:
        lines.append(
            f"#{e['event_number']} | {e['event_date']} ({e['day_name']})\n"
            f"  📍 {e['location'] or '-'} | 🛠️ {e['job'] or '-'}\n"
            f"  👤 معي: {e['partner'] or '-'} | 💵 {fmt_money(e['money'])}\n"
        )
        rows.append(
            [
                InlineKeyboardButton(f"✏️ تعديل #{e['event_number']}", callback_data=f"ef:{e['event_number']}"),
                InlineKeyboardButton(f"🗑️ حذف #{e['event_number']}", callback_data=f"ed:{e['event_number']}:ask"),
            ]
        )

    nav_row = []
    filt_tag = "search" if filters else "all"
    if page > 0:
        nav_row.append(InlineKeyboardButton("⬅️ السابق", callback_data=f"lp:{filt_tag}:{page - 1}"))
    if page < total_pages - 1:
        nav_row.append(InlineKeyboardButton("التالي ➡️", callback_data=f"lp:{filt_tag}:{page + 1}"))
    if nav_row:
        rows.append(nav_row)
    if filters:
        rows.append([InlineKeyboardButton("❌ مسح البحث", callback_data="lp:clear:0")])

    return "\n".join(lines), InlineKeyboardMarkup(rows)


# ----------------------------------------------------------------------
# أوامر عامة
# ----------------------------------------------------------------------

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    ensure_user(user.id, user.full_name or "")
    reset_flow(context)
    await update.message.reply_text(
        "أهلاً بك 👋\n"
        "هذا بوت لتسجيل أيام الشغل مع محمد زيدان (وأي حد بتشتغل معه).\n\n"
        "اختر من القائمة تحت:",
        reply_markup=MAIN_MENU,
    )


async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "الأوامر المتاحة:\n"
        "➕ إضافة يومية - تسجيل يوم شغل جديد\n"
        "📋 آخر الأيام - عرض كل الأيام مع أزرار تعديل/حذف\n"
        "📊 ملخص الحساب - عرض ملخص الحساب الكامل\n"
        "🔍 بحث - فلترة الأيام حسب التاريخ/المكان/الشغل/الشريك\n"
        "📤 تصدير البيانات - Excel أو PDF\n"
        "📈 تقرير شهري/سنوي - رسم بياني للدخل\n"
        "⏰ التذكير الأسبوعي - إرسال ملخص تلقائي كل أسبوع\n"
        "⚙️ الإعدادات - تعديل الرصيد السابق/الخصم/المستلم/الأجرة الافتراضية\n\n"
        "/cancel - لإلغاء أي عملية جارية\n\n"
        "🛠️ Made by OmarMidlg",
        reply_markup=MAIN_MENU,
    )


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    reset_flow(context)
    await update.message.reply_text("تم الإلغاء ❌", reply_markup=MAIN_MENU)


# ----------------------------------------------------------------------
# معالج الأزرار الرئيسية (ReplyKeyboard) + إدخال النصوص أثناء أي عملية جارية
# ----------------------------------------------------------------------

async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    ensure_user(user_id, update.effective_user.full_name or "")
    text = update.message.text.strip()

    awaiting = context.user_data.get("awaiting")
    if awaiting:
        await process_awaiting(update, context, awaiting, text)
        return

    menu_actions = {
        "➕ إضافة يومية": action_add_start,
        "📋 آخر الأيام": action_show_events,
        "📊 ملخص الحساب": action_show_summary,
        "🔍 بحث": action_search_start,
        "📤 تصدير البيانات": action_export_start,
        "📈 تقرير شهري/سنوي": action_report_start,
        "⏰ التذكير الأسبوعي": action_reminder_start,
        "⚙️ الإعدادات": action_settings_start,
    }
    handler = menu_actions.get(text)
    if handler:
        await handler(update, context)
    else:
        await update.message.reply_text("استخدم الأزرار بالأسفل من فضلك 🙂", reply_markup=MAIN_MENU)


# ---------------- إضافة يومية ----------------

async def action_add_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["new_event"] = {}
    context.user_data["awaiting"] = {"type": "add_date"}
    await update.message.reply_text(
        "📅 أرسل تاريخ اليومية بصيغة يوم/شهر/سنة (مثال: 21/09/2026)\n"
        "أو اكتب 'اليوم' لاستخدام تاريخ اليوم." + CANCEL_HINT
    )


async def process_add_date(update, context, text):
    dt = parse_date(text)
    if not dt:
        await update.message.reply_text("صيغة التاريخ غير صحيحة، حاول مثلاً: 21/09/2026 أو اكتب 'اليوم'.")
        return
    day_name = ARABIC_DAYS[dt.weekday()]
    context.user_data["new_event"]["event_date"] = dt.strftime("%d/%m/%Y")
    context.user_data["new_event"]["day_name"] = day_name
    context.user_data["awaiting"] = {"type": "add_location"}
    await update.message.reply_text(f"اليوم: {day_name} ✅\n\n📍 وين كان مكان الحدث؟" + CANCEL_HINT)


async def process_add_location(update, context, text):
    context.user_data["new_event"]["location"] = text
    context.user_data["awaiting"] = {"type": "add_job"}
    await update.message.reply_text("🛠️ شو كان نوع الشغل؟" + CANCEL_HINT)


async def process_add_job(update, context, text):
    context.user_data["new_event"]["job"] = text
    context.user_data["awaiting"] = {"type": "add_partner"}
    await update.message.reply_text("👤 مين اشتغل معك بهاليوم؟ (اكتب '-' إذا اشتغلت لحالك)" + CANCEL_HINT)


async def process_add_partner(update, context, text):
    context.user_data["new_event"]["partner"] = text
    context.user_data["awaiting"] = {"type": "add_money"}
    settings = get_user_settings(update.effective_user.id)
    default_fee = settings["default_fee"]
    await update.message.reply_text(
        f"💵 كم المبلغ لهاليومية؟\n"
        f"أرسل رقم، أو اكتب '=' لاستخدام الأجرة الافتراضية ({fmt_money(default_fee)})." + CANCEL_HINT
    )


async def process_add_money(update, context, text):
    user_id = update.effective_user.id
    if text.strip() == "=":
        settings = get_user_settings(user_id)
        money = settings["default_fee"]
    else:
        money = parse_number(text)

    if money is None or money < 0:
        await update.message.reply_text("الرجاء إرسال رقم صحيح للمبلغ، أو '=' لاستخدام الأجرة الافتراضية.")
        return

    ev = context.user_data["new_event"]
    event_number = get_next_event_number(user_id)
    add_event(
        user_id=user_id,
        event_number=event_number,
        event_date=ev["event_date"],
        day_name=ev["day_name"],
        location=ev["location"],
        job=ev["job"],
        partner=ev["partner"],
        money=money,
    )
    reset_flow(context)

    await update.message.reply_text(
        "✅ تم تسجيل اليومية بنجاح!\n\n"
        f"رقم الحدث: #{event_number}\n"
        f"التاريخ: {ev['event_date']} ({ev['day_name']})\n"
        f"المكان: {ev['location']}\n"
        f"الشغل: {ev['job']}\n"
        f"اشتغل معي: {ev['partner']}\n"
        f"المبلغ: {fmt_money(money)}",
        reply_markup=MAIN_MENU,
    )


# ---------------- عرض الأيام / ملخص ----------------

async def action_show_events(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("current_filters", None)
    text, kb = render_events_page(update.effective_user.id, 0, filters=None)
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=kb or MAIN_MENU)
    if kb:
        await update.message.reply_text("استخدم الأزرار فوق للتنقل أو التعديل/الحذف.", reply_markup=MAIN_MENU)


async def action_show_summary(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = build_summary_text(update.effective_user.id)
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=MAIN_MENU)


# ---------------- بحث ----------------

async def action_search_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    kb = InlineKeyboardMarkup(
        [
            [InlineKeyboardButton("📅 بالتاريخ", callback_data="srch:date")],
            [InlineKeyboardButton("📍 بالمكان", callback_data="srch:location")],
            [InlineKeyboardButton("🛠️ بنوع الشغل", callback_data="srch:job")],
            [InlineKeyboardButton("👤 بالشريك", callback_data="srch:partner")],
            [InlineKeyboardButton("❌ إلغاء", callback_data="srch:cancel")],
        ]
    )
    await update.message.reply_text("🔍 دور بشو بدك تبحث؟", reply_markup=kb)


async def process_search_value(update, context, field, text):
    if field == "date":
        dt = parse_date(text)
        if not dt:
            await update.message.reply_text("صيغة التاريخ غير صحيحة، حاول مثلاً: 21/09/2026 أو 'اليوم'.")
            return
        filters = {"date": dt.strftime("%d/%m/%Y")}
    else:
        filters = {field: text}

    reset_flow(context)
    context.user_data["current_filters"] = filters
    result_text, kb = render_events_page(update.effective_user.id, 0, filters=filters)
    await update.message.reply_text(result_text, parse_mode="Markdown", reply_markup=kb or MAIN_MENU)
    if kb:
        await update.message.reply_text("نتائج البحث ⬆️", reply_markup=MAIN_MENU)


# ---------------- تصدير ----------------

async def action_export_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    kb = InlineKeyboardMarkup(
        [
            [InlineKeyboardButton("📗 Excel", callback_data="exp:xlsx")],
            [InlineKeyboardButton("📕 PDF", callback_data="exp:pdf")],
        ]
    )
    await update.message.reply_text("📤 اختر صيغة التصدير:", reply_markup=kb)


def build_excel_file(user_id) -> io.BytesIO:
    events = get_all_events(user_id)
    settings = get_user_settings(user_id)
    n_events, total_money = get_event_stats(user_id)
    remaining = (
        settings["previous_balance"] + total_money - settings["discount"] - settings["amount_received"]
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "الأيام"
    ws.sheet_view.rightToLeft = True

    headers = ["رقم الحدث", "التاريخ", "اليوم", "المكان", "نوع الشغل", "الشريك", "المبلغ (₪)"]
    ws.append(headers)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for e in events:
        ws.append(
            [e["event_number"], e["event_date"], e["day_name"], e["location"],
             e["job"], e["partner"], e["money"]]
        )

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    # ورقة الملخص
    ws2 = wb.create_sheet("ملخص الحساب")
    ws2.sheet_view.rightToLeft = True
    summary_rows = [
        ("الرصيد من الحساب السابق", settings["previous_balance"]),
        ("عدد الأيام", n_events),
        ("الصافي من الأيام", total_money),
        ("مبلغ الخصم", settings["discount"]),
        ("المبلغ المستلم", settings["amount_received"]),
        ("الرصيد المتبقي", remaining),
    ]
    for label, value in summary_rows:
        ws2.append([label, value])
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 18
    for row in ws2.iter_rows(min_row=1, max_row=len(summary_rows)):
        row[0].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ArabicPDF(FPDF):
    def __init__(self):
        super().__init__(orientation="L", unit="mm", format="A4")
        if os.path.exists(ARABIC_FONT_PATH):
            self.add_font("Arabic", "", ARABIC_FONT_PATH)
            self.font_family_name = "Arabic"
        else:
            self.font_family_name = "Helvetica"

    def ar(self, text):
        return ar_txt(str(text))


def build_pdf_file(user_id) -> io.BytesIO:
    events = get_all_events(user_id)
    settings = get_user_settings(user_id)
    n_events, total_money = get_event_stats(user_id)
    remaining = (
        settings["previous_balance"] + total_money - settings["discount"] - settings["amount_received"]
    )

    pdf = ArabicPDF()
    pdf.add_page()
    pdf.set_font(pdf.font_family_name, size=16)
    pdf.cell(0, 12, pdf.ar("تقرير أيام الشغل"), align="C", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font(pdf.font_family_name, size=10)
    headers = ["رقم", "التاريخ", "اليوم", "المكان", "الشغل", "الشريك", "المبلغ"]
    col_widths = [15, 28, 25, 50, 50, 45, 30]
    # الأعمدة اللي محتواها أرقام/تواريخ (لاتيني بالكامل) تُطبع بالخط الأساسي
    # حتى ما تنكسر رموز متل / و . اللي الخط العربي ما بيدعمها
    numeric_cols = {0, 1, 6}  # رقم، تاريخ، مبلغ

    pdf.set_fill_color(68, 114, 196)
    pdf.set_text_color(255, 255, 255)
    for w, h in zip(col_widths, headers):
        pdf.cell(w, 9, pdf.ar(h), border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_text_color(0, 0, 0)
    fill = False
    for e in events:
        pdf.set_fill_color(240, 240, 240)
        row = [
            e["event_number"], e["event_date"], e["day_name"],
            e["location"] or "-", e["job"] or "-", e["partner"] or "-",
            f"{e['money']:.2f}",
        ]
        for idx, (w, val) in enumerate(zip(col_widths, row)):
            if idx not in numeric_cols and has_arabic(val):
                pdf.set_font(pdf.font_family_name, size=10)
                pdf.cell(w, 8, pdf.ar(val), border=1, align="C", fill=fill)
            else:
                pdf.set_font("helvetica", size=10)
                pdf.cell(w, 8, str(val), border=1, align="C", fill=fill)
        pdf.ln()
        fill = not fill

    pdf.ln(6)
    summary_lines = [
        ("الرصيد من الحساب السابق", f"{settings['previous_balance']:.2f}"),
        ("عدد الأيام", str(n_events)),
        ("الصافي من الأيام", f"{total_money:.2f}"),
        ("مبلغ الخصم", f"{settings['discount']:.2f}"),
        ("المبلغ المستلم", f"{settings['amount_received']:.2f}"),
        ("الرصيد المتبقي", f"{remaining:.2f}"),
    ]
    label_w, value_w = 70, 30
    right_margin = pdf.w - pdf.r_margin
    for label, value in summary_lines:
        pdf.set_x(right_margin - label_w - value_w)
        pdf.set_font(pdf.font_family_name, size=12)
        pdf.cell(label_w, 8, pdf.ar(f"{label}:"), align="R")
        pdf.set_font("helvetica", size=12)
        pdf.cell(value_w, 8, value, align="L", new_x="LMARGIN", new_y="NEXT")

    out = pdf.output()
    return io.BytesIO(bytes(out))


# ---------------- تقارير ورسوم بيانية ----------------

async def action_report_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    kb = InlineKeyboardMarkup(
        [
            [InlineKeyboardButton("آخر 6 أشهر", callback_data="rep:m6")],
            [InlineKeyboardButton("آخر 12 شهر", callback_data="rep:m12")],
            [InlineKeyboardButton("حسب السنة", callback_data="rep:year")],
        ]
    )
    await update.message.reply_text("📈 اختر نوع التقرير:", reply_markup=kb)


def build_bar_chart(labels, values, title):
    fig, ax = plt.subplots(figsize=(8, 4.5))
    display_labels = [ar_txt(str(l)) for l in labels]

    bars = ax.bar(range(len(values)), values, color="#4472C4")
    ax.set_xticks(range(len(display_labels)))
    if ARABIC_FONT_NAME:
        ax.set_xticklabels(display_labels, fontproperties=font_manager.FontProperties(fname=ARABIC_FONT_PATH), rotation=30, ha="right")
        ax.set_title(ar_txt(title), fontproperties=font_manager.FontProperties(fname=ARABIC_FONT_PATH))
    else:
        ax.set_xticklabels(display_labels, rotation=30, ha="right")
        ax.set_title(title)

    for bar, val in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{val:.0f}",
                ha="center", va="bottom", fontsize=8)

    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------------- تذكير أسبوعي ----------------

async def action_reminder_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    settings = get_user_settings(update.effective_user.id)
    status = "مفعّل ✅" if settings["reminder_enabled"] else "متوقف ❌"
    day_label = ARABIC_DAYS.get(settings["reminder_day"], "-")
    kb = InlineKeyboardMarkup(
        [
            [InlineKeyboardButton("✅ تفعيل / تغيير اليوم", callback_data="rem:on")],
            [InlineKeyboardButton("⛔ إيقاف", callback_data="rem:off")],
        ]
    )
    await update.message.reply_text(
        f"⏰ *التذكير الأسبوعي*\n"
        f"الحالة الحالية: {status}\n"
        f"اليوم المحدد: {day_label} الساعة {settings['reminder_hour']}:00\n\n"
        "⚠️ ملاحظة: التذكير بيشتغل فقط إذا كان البوت شغّال بشكل دائم (مستضاف على سيرفر)، مش لما يكون مسكّر على جهازك.",
        parse_mode="Markdown",
        reply_markup=kb,
    )


# ---------------- الإعدادات ----------------

async def action_settings_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    kb = InlineKeyboardMarkup(
        [
            [InlineKeyboardButton("💰 الرصيد من الحساب السابق", callback_data="set:bal")],
            [InlineKeyboardButton("➖ مبلغ الخصم", callback_data="set:disc")],
            [InlineKeyboardButton("✅ المبلغ المستلم", callback_data="set:rec")],
            [InlineKeyboardButton("🏷️ الأجرة الافتراضية لكل حدث", callback_data="set:fee")],
        ]
    )
    await update.message.reply_text("⚙️ اختر ما تريد تعديله:", reply_markup=kb)


async def process_setting_value(update, context, code, text):
    value = parse_number(text)
    if value is None or value < 0:
        await update.message.reply_text("الرجاء إرسال رقم صحيح.")
        return
    field, label = SETTINGS_FIELDS[code]
    update_user_field(update.effective_user.id, field, value)
    reset_flow(context)
    await update.message.reply_text(f"✅ تم تحديث «{label}» إلى {fmt_money(value)}.", reply_markup=MAIN_MENU)


# ---------------- تعديل / حذف يومية (إدخال القيمة الجديدة) ----------------

async def process_edit_value(update, context, event_number, code, text):
    column = FIELD_CODES[code]
    if column == "money":
        value = parse_number(text)
        if value is None or value < 0:
            await update.message.reply_text("الرجاء إرسال رقم صحيح للمبلغ.")
            return
        update_event_field(update.effective_user.id, event_number, "money", value)
    elif column == "event_date":
        dt = parse_date(text)
        if not dt:
            await update.message.reply_text("صيغة التاريخ غير صحيحة، حاول مثلاً: 21/09/2026 أو 'اليوم'.")
            return
        update_event_field(update.effective_user.id, event_number, "event_date", dt.strftime("%d/%m/%Y"))
        update_event_field(update.effective_user.id, event_number, "day_name", ARABIC_DAYS[dt.weekday()])
    else:
        update_event_field(update.effective_user.id, event_number, column, text)

    reset_flow(context)
    await update.message.reply_text(
        f"✅ تم تعديل «{FIELD_LABELS_AR[code]}» لليومية #{event_number}.", reply_markup=MAIN_MENU
    )


# ---------------- موزّع الإدخال النصي أثناء أي عملية جارية ----------------

async def process_awaiting(update: Update, context: ContextTypes.DEFAULT_TYPE, awaiting: dict, text: str):
    t = awaiting["type"]
    if t == "add_date":
        await process_add_date(update, context, text)
    elif t == "add_location":
        await process_add_location(update, context, text)
    elif t == "add_job":
        await process_add_job(update, context, text)
    elif t == "add_partner":
        await process_add_partner(update, context, text)
    elif t == "add_money":
        await process_add_money(update, context, text)
    elif t == "setting_value":
        await process_setting_value(update, context, awaiting["code"], text)
    elif t == "search_value":
        await process_search_value(update, context, awaiting["field"], text)
    elif t == "edit_value":
        await process_edit_value(update, context, awaiting["num"], awaiting["field"], text)
    else:
        reset_flow(context)
        await update.message.reply_text("صار في خطأ، جرّب من جديد.", reply_markup=MAIN_MENU)


# ----------------------------------------------------------------------
# معالج الأزرار التفاعلية (Inline)
# ----------------------------------------------------------------------

async def callback_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    data = query.data
    user_id = update.effective_user.id
    ensure_user(user_id, update.effective_user.full_name or "")
    await query.answer()

    parts = data.split(":")
    action = parts[0]

    # ------- ترقيم صفحات قائمة الأيام -------
    if action == "lp":
        mode, page = parts[1], int(parts[2])
        if mode == "clear":
            context.user_data.pop("current_filters", None)
            filters = None
        else:
            filters = context.user_data.get("current_filters")
        text, kb = render_events_page(user_id, page, filters=filters)
        await query.edit_message_text(text, parse_mode="Markdown", reply_markup=kb)
        return

    # ------- فتح قائمة تعديل حقل يومية -------
    if action == "ef":
        event_number = int(parts[1])
        kb = InlineKeyboardMarkup(
            [[InlineKeyboardButton(FIELD_LABELS_AR[c], callback_data=f"efc:{event_number}:{c}")] for c in FIELD_CODES]
            + [[InlineKeyboardButton("❌ إلغاء", callback_data="efc:cancel")]]
        )
        await query.message.reply_text(f"✏️ شو بدك تعدّل باليومية #{event_number}؟", reply_markup=kb)
        return

    if action == "efc":
        if parts[1] == "cancel":
            await query.message.reply_text("تم الإلغاء ❌", reply_markup=MAIN_MENU)
            return
        event_number, code = int(parts[1]), parts[2]
        context.user_data["awaiting"] = {"type": "edit_value", "num": event_number, "field": code}
        await query.message.reply_text(
            f"أرسل القيمة الجديدة لـ «{FIELD_LABELS_AR[code]}» لليومية #{event_number}:" + CANCEL_HINT
        )
        return

    # ------- حذف يومية -------
    if action == "ed":
        event_number, step = int(parts[1]), parts[2]
        if step == "ask":
            kb = InlineKeyboardMarkup(
                [
                    [
                        InlineKeyboardButton("✅ نعم، احذف", callback_data=f"ed:{event_number}:yes"),
                        InlineKeyboardButton("❌ لا", callback_data=f"ed:{event_number}:no"),
                    ]
                ]
            )
            await query.message.reply_text(f"متأكد إنك بدك تحذف اليومية #{event_number}؟", reply_markup=kb)
        elif step == "yes":
            ok = delete_event(user_id, event_number)
            msg = f"✅ تم حذف اليومية #{event_number}." if ok else "❌ ما لقيت هاي اليومية."
            await query.message.reply_text(msg, reply_markup=MAIN_MENU)
        else:
            await query.message.reply_text("تم الإلغاء ❌", reply_markup=MAIN_MENU)
        return

    # ------- بحث -------
    if action == "srch":
        field = parts[1]
        if field == "cancel":
            await query.message.reply_text("تم الإلغاء ❌", reply_markup=MAIN_MENU)
            return
        if field == "date":
            prompt = "📅 أرسل التاريخ (يوم/شهر/سنة) أو 'اليوم':"
        elif field == "location":
            prompt = "📍 اكتب جزء من اسم المكان:"
        elif field == "job":
            prompt = "🛠️ اكتب جزء من نوع الشغل:"
        else:
            prompt = "👤 اكتب جزء من اسم الشريك:"
        context.user_data["awaiting"] = {"type": "search_value", "field": field}
        await query.message.reply_text(prompt + CANCEL_HINT)
        return

    # ------- تصدير -------
    if action == "exp":
        fmt = parts[1]
        n_events, _ = get_event_stats(user_id)
        if n_events == 0:
            await query.message.reply_text("لا يوجد بيانات لتصديرها بعد.", reply_markup=MAIN_MENU)
            return
        await query.message.reply_text("⏳ جارِ تجهيز الملف...")
        if fmt == "xlsx":
            buf = build_excel_file(user_id)
            await context.bot.send_document(
                chat_id=user_id, document=buf, filename="gig_tracker.xlsx", caption="📗 ملف Excel - Made by OmarMidlg"
            )
        else:
            buf = build_pdf_file(user_id)
            await context.bot.send_document(
                chat_id=user_id, document=buf, filename="gig_tracker.pdf", caption="📕 ملف PDF - Made by OmarMidlg"
            )
        return

    # ------- تقارير -------
    if action == "rep":
        kind = parts[1]
        n_events, _ = get_event_stats(user_id)
        if n_events == 0:
            await query.message.reply_text("لا يوجد بيانات كافية لعمل تقرير بعد.", reply_markup=MAIN_MENU)
            return
        await query.message.reply_text("⏳ جارِ تجهيز الرسم البياني...")
        if kind == "m6":
            data = get_monthly_totals(user_id, months_back=6)
            labels = [d[0] for d in data]
            values = [d[1] for d in data]
            title = "الدخل خلال آخر 6 أشهر"
        elif kind == "m12":
            data = get_monthly_totals(user_id, months_back=12)
            labels = [d[0] for d in data]
            values = [d[1] for d in data]
            title = "الدخل خلال آخر 12 شهر"
        else:
            data = get_yearly_totals(user_id)
            labels = [str(d[0]) for d in data]
            values = [d[1] for d in data]
            title = "الدخل حسب السنة"

        chart_buf = build_bar_chart(labels, values, title)
        total = sum(values)
        caption = f"📈 {title}\nالمجموع: {fmt_money(total)}\n\nMade by OmarMidlg"
        await context.bot.send_photo(chat_id=user_id, photo=chart_buf, caption=caption)
        return

    # ------- تذكير أسبوعي -------
    if action == "rem":
        if parts[1] == "off":
            set_reminder(user_id, enabled=False)
            await query.message.reply_text("⛔ تم إيقاف التذكير الأسبوعي.", reply_markup=MAIN_MENU)
            return
        if parts[1] == "on":
            kb = InlineKeyboardMarkup(
                [[InlineKeyboardButton(ARABIC_DAYS[d], callback_data=f"remday:{d}")] for d in ARABIC_DAY_ORDER]
            )
            await query.message.reply_text("اختر يوم الأسبوع لإرسال التذكير (الساعة 18:00):", reply_markup=kb)
            return

    if action == "remday":
        day = int(parts[1])
        set_reminder(user_id, enabled=True, day=day, hour=18)
        await query.message.reply_text(
            f"✅ تم تفعيل التذكير الأسبوعي: كل {ARABIC_DAYS[day]} الساعة 18:00.\n"
            "⚠️ تذكّر إنه لازم البوت يكون مستضاف ومشغّل بشكل دائم حتى يوصلك التذكير.",
            reply_markup=MAIN_MENU,
        )
        return

    # ------- الإعدادات -------
    if action == "set":
        code = parts[1]
        field, label = SETTINGS_FIELDS[code]
        settings = get_user_settings(user_id)
        current_value = settings[field]
        context.user_data["awaiting"] = {"type": "setting_value", "code": code}
        await query.message.reply_text(
            f"القيمة الحالية لـ «{label}»: {fmt_money(current_value)}\n\n"
            f"أرسل القيمة الجديدة (رقم فقط):" + CANCEL_HINT
        )
        return


# ----------------------------------------------------------------------
# التذكير الأسبوعي التلقائي (Job Queue)
# ----------------------------------------------------------------------

async def reminder_job(context: ContextTypes.DEFAULT_TYPE):
    now = datetime.now()
    today_str = now.strftime("%Y-%m-%d")
    users = get_all_users_with_reminders()
    for u in users:
        if u["reminder_day"] != now.weekday():
            continue
        if u["reminder_hour"] != now.hour:
            continue
        if u["last_reminder_sent"] == today_str:
            continue
        try:
            text = "⏰ *تذكير أسبوعي*\n\n" + build_summary_text(u["user_id"])
            await context.bot.send_message(chat_id=u["user_id"], text=text, parse_mode="Markdown")
            set_last_reminder_sent(u["user_id"], today_str)
        except Exception as exc:  # noqa: BLE001
            logger.warning("تعذّر إرسال تذكير للمستخدم %s: %s", u["user_id"], exc)


# ----------------------------------------------------------------------
# تشغيل البوت
# ----------------------------------------------------------------------

def main():
    if not BOT_TOKEN:
        raise SystemExit("الرجاء ضبط متغير BOT_TOKEN بتوكن البوت الخاص بك قبل التشغيل.")

    init_db()

    app = Application.builder().token(BOT_TOKEN).build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(CommandHandler("cancel", cancel))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_handler(CallbackQueryHandler(callback_router))

    if app.job_queue is not None:
        app.job_queue.run_repeating(reminder_job, interval=3600, first=15)
    else:
        logger.warning("JobQueue غير متوفر - التذكير الأسبوعي التلقائي لن يعمل. تأكد من تثبيت: pip install \"python-telegram-bot[job-queue]\"")

    # حل مشكلة Python 3.14 التي أزالت الإنشاء الضمني لحلقة الأحداث بالخيط الرئيسي
    import asyncio
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        try:
            asyncio.get_event_loop()
        except RuntimeError:
            asyncio.set_event_loop(asyncio.new_event_loop())

    logger.info("Bot starting... Made by OmarMidlg")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
